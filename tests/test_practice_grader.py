"""
Pruebas del corrector automático de prácticas.

Excel, al guardar, deja junto a cada fórmula su resultado calculado (<v>).
Aquí no hay Excel, así que `excel_save()` simula eso: escribe las fórmulas del
"alumno" con openpyxl y luego inyecta en el XML los resultados que Excel habría
guardado.

Correr:  python -m pytest tests/test_practice_grader.py -q
"""
import io
import re
import zipfile
from xml.sax.saxutils import escape

import pytest
from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName

from app.python import excel_generator, excel_generator_en
from app.python.practice_grader import grade_practice


def practice(level, lang="es"):
    gen = excel_generator if lang == "es" else excel_generator_en
    return load_workbook(io.BytesIO(gen.generate_excel(level, "practica", "en")))


def excel_save(wb, answers, cached=True, sheet=None):
    """answers: {"B21": ("=SUM(G6:G15)", 14310), "C5": ("texto", None) ...}
    Valor tupla (formula_o_valor, resultado_cacheado). Para errores: resultado "#N/A"."""
    ws = wb[sheet] if sheet else wb.worksheets[0]
    for coord, (val, _) in answers.items():
        ws[coord] = val
    buf = io.BytesIO()
    wb.save(buf)
    if not cached:
        return buf.getvalue()
    idx = wb.worksheets.index(ws) + 1
    return _inject(buf.getvalue(), f"xl/worksheets/sheet{idx}.xml",
                   {k: v for k, (f, v) in answers.items() if isinstance(f, str) and f.startswith("=")})


def _inject(data, sheet_path, values):
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            raw = zin.read(info.filename)
            if info.filename == sheet_path:
                xml = raw.decode("utf-8")
                for coord, val in values.items():
                    xml = _set_cached(xml, coord, val)
                raw = xml.encode("utf-8")
            zout.writestr(info, raw)
    return out.getvalue()


def _set_cached(xml, coord, val):
    pat = re.compile(r'<c r="%s"([^>]*)>(.*?)</c>' % coord, re.S)
    m = pat.search(xml)
    assert m, coord
    attrs = re.sub(r'\s+t="[^"]*"', "", m.group(1))
    inner = re.sub(r"<v\s*/>|<v>.*?</v>", "", m.group(2), flags=re.S)
    if isinstance(val, str) and val.startswith("#"):
        attrs += ' t="e"'
        v = escape(val)
    elif isinstance(val, str):
        attrs += ' t="str"'
        v = escape(val)
    elif isinstance(val, bool):
        attrs += ' t="b"'
        v = "1" if val else "0"
    else:
        v = repr(val)
    return xml[:m.start()] + f'<c r="{coord}"{attrs}>{inner}<v>{v}</v></c>' + xml[m.end():]


def statuses(res):
    return [it["status"] for ex in res["exercises"] for it in ex["items"]]


# ── Nivel 1 ──────────────────────────────────────────────────────────────────

L1_TOTALS = [2550, 360, 640, 600, 3400, 450, 960, 540, 560, 4250]


def l1_perfect():
    a = {
        "B21": ("=SUM(G6:G15)", 14310), "B22": ("=AVERAGE(G6:G15)", 1431),
        "B23": ("=COUNT(G6:G15)", 10), "B24": ("=MAX(G6:G15)", 4250),
        "B25": ("=MIN(G6:G15)", 360), "B26": ("=ROUND(SUM(G6:G15),-1)", 14310),
    }
    for i, g in enumerate(L1_TOTALS):
        a[f"C{32 + i}"] = (f"=IF(G{6 + i}>2000,G{6 + i}*5%,G{6 + i}*2%)", g * (0.05 if g > 2000 else 0.02))
    return a


@pytest.mark.parametrize("lang", ["es", "en"])
def test_l1_perfect(lang):
    res = grade_practice(excel_save(practice("nivel1", lang), l1_perfect()), "nivel1", "x.xlsx", lang)
    assert res["status"] == "graded" and res["score"] == 100, res
    assert not res["warnings"]


def test_l1_hardcoded_wrong_and_errors():
    a = l1_perfect()
    a["B21"] = (14310, None)                      # escrito a mano
    a["B22"] = ("=SUM(G6:G15)/9", 1590)           # resultado incorrecto
    a["B23"] = ("=COUNT(G6:G15", "#NAME?")        # error
    a["C32"] = ("=G6*5%", 127.5)                  # sin IF y resultado correcto → parcial
    res = grade_practice(excel_save(practice("nivel1"), a), "nivel1", "x.xlsx")
    st = statuses(res)
    assert st[0] == "fail" and "mano" in res["exercises"][0]["items"][0]["msg"]["es"]
    assert st[1] == "fail" and st[2] == "fail" and "#NAME?" in res["exercises"][0]["items"][2]["msg"]["es"]
    assert st[6] == "partial"
    assert 0 < res["score"] < 100


def test_l1_not_saved_in_excel_is_unverified():
    res = grade_practice(excel_save(practice("nivel1"), l1_perfect(), cached=False), "nivel1", "x.xlsx")
    assert set(statuses(res)) == {"unverified"}
    assert res["score"] == 0 and any("Excel" in w["es"] for w in res["warnings"])


def test_l1_inserted_rows_renamed_sheet_and_answer_next_column():
    wb = practice("nivel1")
    ws = wb.worksheets[0]
    ws.title = "Mi tarea"
    ws.insert_rows(1, 3)  # el alumno metió 3 filas arriba
    a = {}
    for k, (f, v) in l1_perfect().items():
        col, row = re.match(r"([A-Z]+)(\d+)", k).groups()
        f = re.sub(r"G(\d+)", lambda m: f"G{int(m.group(1)) + 3}", f)
        a[f"{col}{int(row) + 3}"] = (f, v)
    a["F24"] = a.pop("B24")  # respuesta escrita en otra columna de la misma fila
    res = grade_practice(excel_save(wb, a), "nivel1", "x.xlsx")
    assert res["score"] == 100, [(it["label"], it["status"], it["msg"]) for ex in res["exercises"] for it in ex["items"]]


def test_l1_changed_data_graded_against_own_data():
    wb = practice("nivel1")
    wb.worksheets[0]["E6"] = 10  # 10 * 850 = 8500 en vez de 2550
    a = l1_perfect()
    a["B21"] = ("=SUM(G6:G15)", 14310 - 2550 + 8500)
    res = grade_practice(excel_save(wb, a), "nivel1", "x.xlsx")
    assert res["exercises"][0]["items"][0]["status"] == "ok"
    assert res["warnings"]


def test_l1_commission_in_column_h():
    a = {k: v for k, v in l1_perfect().items() if not k.startswith("C")}
    for i, g in enumerate(L1_TOTALS):
        a[f"H{6 + i}"] = (f"=IF(G{6 + i}>2000,G{6 + i}*0.05,G{6 + i}*0.02)", g * (0.05 if g > 2000 else 0.02))
    res = grade_practice(excel_save(practice("nivel1"), a), "nivel1", "x.xlsx")
    assert res["score"] == 100


# ── Nivel 0 ──────────────────────────────────────────────────────────────────

def test_l0():
    from openpyxl.worksheet.table import Table
    wb = practice("nivel0")
    ws = wb.worksheets[0]
    ws.add_table(Table(displayName="Contactos", ref="A6:E11"))
    ws.column_dimensions["C"].width = 24
    answers = ["Texto", "número", "Numero", "FECHA", "Booleano (lógico)", "error", "Text", "Fórmula"]
    a = {f"B{17 + i}": (ans, None) for i, ans in enumerate(answers)}
    res = grade_practice(excel_save(wb, a), "nivel0", "x.xlsx")
    assert res["score"] == 100, [(it["label"], it["status"]) for ex in res["exercises"] for it in ex["items"]]
    a["B18"] = ("Texto", None)
    a["B19"] = ("Texto o número", None)
    res = grade_practice(excel_save(practice("nivel0"), a), "nivel0", "x.xlsx")
    st = statuses(res)
    assert st[:2] == ["fail", "fail"]  # sin tabla ni anchos
    assert st[3] == "fail" and st[4] == "fail"


# ── Nivel 2 ──────────────────────────────────────────────────────────────────

@pytest.mark.parametrize("lang", ["es", "en"])
def test_l2(lang):
    dom = "@empresa.com" if lang == "es" else "@company.com"
    sums = [8800, 9000, 2900, 3600]
    a = {f"B{19 + i}": (f"=SUMIFS($E$6:$E$13,$C$6:$C$13,A{19 + i})", s) for i, s in enumerate(sums)}
    locals_es = ["agarcía", "clópez", "mtorres", "lmorales", "schen"]
    locals_en = ["agreen", "clewis", "mturner", "lmoore", "schen"]
    for i, l in enumerate(locals_es if lang == "es" else locals_en):
        a[f"B{28 + i}"] = (f'=LOWER(LEFT(A{28 + i},1)&MID(A{28 + i},FIND(" ",A{28 + i})+1,100)&"{dom}")', l + dom)
    res = grade_practice(excel_save(practice("nivel2", lang), a), "nivel2", "x.xlsx", lang)
    assert res["score"] == 100, [(it["label"], it["status"], it["got"], it["expected"]) for ex in res["exercises"] for it in ex["items"]]
    a["B19"] = ("=SUMIF($C$6:$C$13,A19,$E$6:$E$13)", 8800)  # SUMIF: medio punto
    a["B28"] = ('=LEFT(A28,1)&MID(A28,FIND(" ",A28)+1,100)&"' + dom + '"', "AGarcia" + dom)  # sin LOWER pero se acepta? no: mayúsculas
    res = grade_practice(excel_save(practice("nivel2", lang), a), "nivel2", "x.xlsx", lang)
    st = statuses(res)
    assert st[0] == "partial"
    assert st[4] == "fail"  # email con mayúsculas: el ejercicio pide minúsculas


# ── Nivel 3 ──────────────────────────────────────────────────────────────────

def l3_answers(code="P-005"):
    recs = {"P-005": ["Fuente 850W", "PSU", 75, 149, 15, "OK"], "P-009": ["SSD NVMe 2TB", "Almacenamiento", 90, 179, 1, "🔴 CRÍTICO"]}
    letters = ["B", "C", "D", "E", "F", "H"]
    a = {"C20": (code, None)}
    for i, (L, v) in enumerate(zip(letters, recs[code])):
        a[f"C{23 + i}"] = (f'=XLOOKUP($C$20,$A$6:$A$15,${L}$6:${L}$15,"No encontrado")', v)
    return a


def test_l3_lookup_and_filter():
    wb = practice("nivel3")
    ws = wb.worksheets[0]
    a = l3_answers("P-009")
    a["B33"] = ("=FILTER($A$6:$H$15,$F$6:$F$15<$G$6:$G$15,\"Sin productos\")", "P-004")
    # celdas desbordadas que Excel guarda como valores
    ws["B34"] = None
    ws["B34"] = "P-006"
    ws["B35"] = "P-009"
    res = grade_practice(excel_save(wb, a), "nivel3", "x.xlsx")
    assert res["score"] == 100, [(it["label"], it["status"], it["msg"], it["got"]) for ex in res["exercises"] for it in ex["items"]]


def test_l3_direct_reference_and_spill_from_old_layout():
    wb = practice("nivel3")
    a = l3_answers()
    a["C23"] = ("=B10", "Fuente 850W")  # referencia directa
    a["B33"] = ("=FILTER(A6:H15,F6:F15<G6:G15)", "#SPILL!")
    res = grade_practice(excel_save(wb, a), "nivel3", "x.xlsx")
    it0 = res["exercises"][0]["items"][0]
    assert it0["status"] == "partial" and "referencia directa" in it0["msg"]["es"]
    assert res["exercises"][1]["items"][0]["status"] == "ok"


def test_l3_wrong_filter_condition():
    wb = practice("nivel3")
    a = l3_answers()
    a["B33"] = ("=FILTER(A6:H15,F6:F15>G6:G15)", "P-001")
    res = grade_practice(excel_save(wb, a), "nivel3", "x.xlsx")
    assert res["exercises"][1]["items"][0]["status"] == "fail"


# ── Nivel 4 ──────────────────────────────────────────────────────────────────

def test_l4():
    wb = practice("nivel4")
    for name, txt in [("IVA", "_xlfn.LAMBDA(_xlpm.monto,_xlpm.monto*0.16)"),
                      ("PRECIO_FINAL", "_xlfn.LAMBDA(_xlpm.p,_xlpm.d,_xlpm.p*(1-_xlpm.d))"),
                      ("DIAS_HABILES", "_xlfn.LAMBDA(_xlpm.i,_xlpm.f,NETWORKDAYS(_xlpm.i,_xlpm.f))"),
                      ("CATEGORIZAR", '_xlfn.LAMBDA(_xlpm.n,IF(_xlpm.n>=9,"A",IF(_xlpm.n>=7,"B","C")))')]:
        wb.defined_names[name] = DefinedName(name, attr_text=txt)
    exp = ["⭐ Cliente A", "🟡 Cliente B", "Cliente C", "⭐ Cliente A", "🟡 Cliente B"]
    a = {f"C{26 + i}": (f'=IFS(B{26 + i}>5000,"⭐ Cliente A",B{26 + i}>=2000,"🟡 Cliente B",TRUE,"Cliente C")', e)
         for i, e in enumerate(exp)}
    res = grade_practice(excel_save(wb, a), "nivel4", "x.xlsx")
    st = statuses(res)
    assert st[:4] == ["ok"] * 4
    assert st[-5:] == ["ok"] * 5
    assert st[4:9] == ["fail", "fail", "fail", "fail", "fail"]  # dashboard no construido


# ── Archivos raros ───────────────────────────────────────────────────────────

def test_garbage_and_wrong_files():
    assert grade_practice(b"no soy un excel", "nivel1", "x.xlsx")["status"] == "unreadable"
    assert grade_practice(b"\xd0\xcf\x11\xe0", "nivel1", "viejo.xls")["status"] == "unreadable"
    from openpyxl import Workbook
    buf = io.BytesIO()
    Workbook().save(buf)
    assert grade_practice(buf.getvalue(), "nivel1", "x.xlsx")["status"] == "wrong_file"
    # entregar la práctica en blanco no revienta
    res = grade_practice(excel_save(practice("nivel2"), {}), "nivel2", "x.xlsx")
    assert res["score"] == 0 and set(statuses(res)) == {"fail"}


def test_formula_typed_as_text():
    a = l1_perfect()
    a["B21"] = ("=SUM(G6:G15)", None)
    wb = practice("nivel1")
    wb.worksheets[0]["B21"].number_format = "@"
    data = excel_save(wb, {k: v for k, v in a.items() if k != "B21"})
    # Simula una celda en formato Texto: Excel la guarda como cadena, no como fórmula
    wb2 = load_workbook(io.BytesIO(data))
    wb2.worksheets[0]["B21"].value = "=SUM(G6:G15)"
    wb2.worksheets[0]["B21"].data_type = "s"
    buf = io.BytesIO(); wb2.save(buf)
    xml_fix = buf.getvalue()
    res = grade_practice(xml_fix, "nivel1", "x.xlsx")
    it = res["exercises"][0]["items"][0]
    assert it["status"] == "fail" and "texto" in it["msg"]["es"], it
