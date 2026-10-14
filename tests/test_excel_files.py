"""
Los .xlsx descargables deben abrir sin #¿NOMBRE? en Excel de CUALQUIER idioma.

Un .xlsx guarda las fórmulas en inglés (Excel las traduce al abrir). Esta prueba
genera todos los archivos (5 niveles × ejemplos/práctica × preferencia es/en ×
generador ES/EN) y falla si alguna celda-fórmula usa una función que no existe
en Excel en inglés, o si hay texto guardado como fórmula rota.
"""
import io
import re

import pytest
from openpyxl import load_workbook

from app.python import excel_generator, excel_generator_en

VALID = set("""ABS AGGREGATE AND AVERAGE AVERAGEIF AVERAGEIFS CHOOSE CONCAT CONCATENATE COUNT COUNTA
COUNTBLANK COUNTIF COUNTIFS DATE DATEDIF DAY DAYS EDATE EOMONTH EXACT FILTER FIND HLOOKUP HOUR IF IFERROR
IFNA IFS INDEX INDIRECT INT ISBLANK ISERROR ISNUMBER ISTEXT LAMBDA LARGE LEFT LEN LET LOOKUP LOWER MATCH
MAX MAXIFS MEDIAN MID MIN MINIFS MINUTE MOD MONTH NETWORKDAYS NOT NOW OFFSET OR PROPER RANK RIGHT ROUND
ROUNDDOWN ROUNDUP ROW ROWS SEARCH SEQUENCE SMALL SORT SORTBY SUBSTITUTE SUBTOTAL SUM SUMIF SUMIFS SUMPRODUCT
SWITCH TEXT TEXTJOIN TEXTBEFORE TEXTAFTER TODAY TRIM UNIQUE UPPER VALUE VLOOKUP WEEKDAY WORKDAY XLOOKUP
XMATCH YEAR TAKE DROP VSTACK HSTACK BYROW MAP REDUCE COLUMN COLUMNS TRANSPOSE""".split())
FUNC = re.compile(r"(?<![A-Za-z0-9_.])([A-Za-zÁÉÍÓÚÑáéíóúñ][A-Za-zÁÉÍÓÚÑáéíóúñ0-9_.]*)\(")

CASES = [(g, pref, lv, t)
         for g in (excel_generator, excel_generator_en)
         for pref in ("es", "en")
         for lv in range(5)
         for t in ("ejemplos", "practica")]


@pytest.mark.parametrize("gen,pref,lv,kind", CASES,
                         ids=[f"{g.__name__.split('.')[-1]}-{p}-nivel{l}-{k}" for g, p, l, k in CASES])
def test_formulas_valid_in_any_excel(gen, pref, lv, kind):
    wb = load_workbook(io.BytesIO(gen.generate_excel(f"nivel{lv}", kind, pref)))
    problems = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type != "f" or not isinstance(c.value, str):
                    continue
                body = re.sub(r'"[^"]*"', '""', c.value)
                for fn in FUNC.findall(body):
                    if fn.upper() not in VALID:
                        problems.append(f"{ws.title}!{c.coordinate}: {fn} en {c.value}")
                if "←" in body or body.count("(") != body.count(")"):
                    problems.append(f"{ws.title}!{c.coordinate}: fórmula rota {c.value}")
    assert not problems, "\n".join(problems[:20])
